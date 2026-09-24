"""
Convert docs/project_report.md to Nivetha_ProjectReport.docx
"""

import os
import re
import docx
from docx.shared import Inches, Pt, RGBColor
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.enum.table import WD_TABLE_ALIGNMENT, WD_ALIGN_VERTICAL
from docx.oxml import parse_xml
from docx.oxml.ns import nsdecls

def create_report_docx(md_path: str, output_path: str):
    with open(md_path, "r", encoding="utf-8") as f:
        md_content = f.read()

    doc = docx.Document()

    # Set page margins
    sections = doc.sections
    for section in sections:
        section.top_margin = Inches(1.0)
        section.bottom_margin = Inches(1.0)
        section.left_margin = Inches(1.0)
        section.right_margin = Inches(1.0)

    # Styles
    normal_style = doc.styles['Normal']
    normal_style.font.name = 'Calibri'
    normal_style.font.size = Pt(11)
    normal_style.font.color.rgb = RGBColor(33, 37, 41)

    lines = md_content.split('\n')
    i = 0
    in_code_block = False
    code_lines = []

    while i < len(lines):
        line = lines[i]

        # Code block handling
        if line.strip().startswith('```'):
            if in_code_block:
                # Flush code block
                code_text = '\n'.join(code_lines)
                table = doc.add_table(rows=1, cols=1)
                table.alignment = WD_TABLE_ALIGNMENT.CENTER
                cell = table.cell(0, 0)
                shd = parse_xml(f'<w:shd {nsdecls("w")} w:fill="F1F5F9"/>')
                cell._tc.get_or_add_tcPr().append(shd)
                p = cell.paragraphs[0]
                p.paragraph_format.space_before = Pt(4)
                p.paragraph_format.space_after = Pt(4)
                run = p.add_run(code_text)
                run.font.name = 'Courier New'
                run.font.size = Pt(9.5)
                run.font.color.rgb = RGBColor(30, 41, 59)
                in_code_block = False
                code_lines = []
            else:
                in_code_block = True
                code_lines = []
            i += 1
            continue

        if in_code_block:
            code_lines.append(line)
            i += 1
            continue

        # Table handling (Markdown tables)
        if line.strip().startswith('|') and '|' in line[1:]:
            table_lines = []
            while i < len(lines) and lines[i].strip().startswith('|'):
                table_lines.append(lines[i].strip())
                i += 1
            
            # Filter out separator line | :--- | :--- |
            cleaned_rows = [r for r in table_lines if not re.match(r'^\|[\s\-:]+(\|[\s\-:]+)+\|$', r)]
            if cleaned_rows:
                headers = [c.strip() for c in cleaned_rows[0].strip('|').split('|')]
                data_rows = [[c.strip() for c in r.strip('|').split('|')] for r in cleaned_rows[1:]]

                t = doc.add_table(rows=len(cleaned_rows), cols=len(headers))
                t.alignment = WD_TABLE_ALIGNMENT.CENTER
                
                # Header row
                hdr_cells = t.rows[0].cells
                for idx, text in enumerate(headers):
                    hdr_cells[idx].text = text
                    shd = parse_xml(f'<w:shd {nsdecls("w")} w:fill="1E293B"/>')
                    hdr_cells[idx]._tc.get_or_add_tcPr().append(shd)
                    p = hdr_cells[idx].paragraphs[0]
                    p.runs[0].font.bold = True
                    p.runs[0].font.color.rgb = RGBColor(255, 255, 255)
                    p.runs[0].font.size = Pt(10)

                # Data rows
                for r_idx, row in enumerate(data_rows):
                    row_cells = t.rows[r_idx + 1].cells
                    fill_color = "F8FAFC" if r_idx % 2 == 1 else "FFFFFF"
                    for c_idx, val in enumerate(row):
                        if c_idx < len(row_cells):
                            row_cells[c_idx].text = val
                            shd = parse_xml(f'<w:shd {nsdecls("w")} w:fill="{fill_color}"/>')
                            row_cells[c_idx]._tc.get_or_add_tcPr().append(shd)
                            p = row_cells[c_idx].paragraphs[0]
                            if p.runs:
                                p.runs[0].font.size = Pt(9.5)
            continue

        # Horizontal rule
        if line.strip() in ['---', '***', '___']:
            p = doc.add_paragraph()
            p.paragraph_format.space_before = Pt(6)
            p.paragraph_format.space_after = Pt(6)
            run = p.add_run('—' * 55)
            run.font.color.rgb = RGBColor(203, 213, 225)
            i += 1
            continue

        # Headings
        if line.startswith('# '):
            p = doc.add_paragraph()
            p.paragraph_format.space_before = Pt(16)
            p.paragraph_format.space_after = Pt(6)
            run = p.add_run(line[2:].strip())
            run.font.name = 'Calibri'
            run.font.size = Pt(22)
            run.font.bold = True
            run.font.color.rgb = RGBColor(15, 23, 42)
        elif line.startswith('## '):
            p = doc.add_paragraph()
            p.paragraph_format.space_before = Pt(14)
            p.paragraph_format.space_after = Pt(4)
            run = p.add_run(line[3:].strip())
            run.font.name = 'Calibri'
            run.font.size = Pt(16)
            run.font.bold = True
            run.font.color.rgb = RGBColor(30, 41, 59)
        elif line.startswith('### '):
            p = doc.add_paragraph()
            p.paragraph_format.space_before = Pt(10)
            p.paragraph_format.space_after = Pt(2)
            run = p.add_run(line[4:].strip())
            run.font.name = 'Calibri'
            run.font.size = Pt(13)
            run.font.bold = True
            run.font.color.rgb = RGBColor(51, 65, 85)
        elif line.startswith('- ') or line.startswith('* '):
            p = doc.add_paragraph(style='List Bullet')
            p.paragraph_format.space_after = Pt(2)
            _add_formatted_runs(p, line[2:].strip())
        elif re.match(r'^\d+\.\s', line):
            p = doc.add_paragraph(style='List Number')
            p.paragraph_format.space_after = Pt(2)
            text = re.sub(r'^\d+\.\s', '', line).strip()
            _add_formatted_runs(p, text)
        elif line.strip():
            p = doc.add_paragraph()
            p.paragraph_format.space_after = Pt(4)
            p.paragraph_format.line_spacing = 1.15
            _add_formatted_runs(p, line.strip())

        i += 1

    doc.save(output_path)
    print(f"Successfully generated {output_path}")


def _add_formatted_runs(paragraph, text: str):
    """Parses markdown bold, italic, inline code, and links to runs."""
    # Simple regex parsing for **bold** and *italic* and `code`
    parts = re.split(r'(\*\*.*?\*\*|\*.*?\*|`.*?`)', text)
    for part in parts:
        if not part:
            continue
        if part.startswith('**') and part.endswith('**'):
            run = paragraph.add_run(part[2:-2])
            run.bold = True
        elif part.startswith('*') and part.endswith('*'):
            run = paragraph.add_run(part[1:-1])
            run.italic = True
        elif part.startswith('`') and part.endswith('`'):
            run = paragraph.add_run(part[1:-1])
            run.font.name = 'Courier New'
            run.font.size = Pt(10)
            run.font.color.rgb = RGBColor(194, 65, 12)
        else:
            # Clean markdown link [text](url) -> text
            cleaned_text = re.sub(r'\[([^\]]+)\]\([^\)]+\)', r'\1', part)
            paragraph.add_run(cleaned_text)


if __name__ == "__main__":
    report_md = "docs/project_report.md"
    report_docx = "Nivetha_ProjectReport.docx"
    create_report_docx(report_md, report_docx)
